"""Documento editable sobrio, coherente con Times New Roman de ACA2."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
HERE=Path(__file__).resolve().parents[1]
def build():
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
    styles=doc.styles
    for node in styles.element.iter(qn('w:pBdr')):
        node.getparent().remove(node)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        styles[name].font.name='Times New Roman';styles[name].font.size=Pt(12);styles[name].font.color.rgb=RGBColor(0,0,0)
        rf=styles[name].element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(rf.attrib):
            if 'Theme' in attr or 'theme' in attr:del rf.attrib[attr]
        rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
    styles['Normal'].paragraph_format.line_spacing=1.5
    styles['Normal'].paragraph_format.first_line_indent=Inches(.5)
    styles['Normal'].paragraph_format.space_after=Pt(6)
    for name in ['Heading 1','Heading 2','Heading 3']:
        styles[name].font.bold=True;styles[name].paragraph_format.keep_with_next=True
        styles[name].paragraph_format.space_before=Pt(12);styles[name].paragraph_format.space_after=Pt(6)
    styles['Title'].font.bold=True;styles['Title'].font.size=Pt(14)
    lines=(HERE/'objetivo3/TEXTO_PARA_TESIS.md').read_text().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):doc.add_paragraph(line[2:],style='Title');continue
        if line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1');continue
        if line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2');continue
        if line.startswith('!['):
            target=re.search(r'\]\((.*?)\)',line).group(1)
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=0;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(HERE/'objetivo3'/target),width=Inches(6.3));continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
            vals=[[v.strip() for v in r.strip('|').split('|')] for r in rows if not re.match(r'^\|\s*---',r)]
            table=doc.add_table(rows=0,cols=len(vals[0]));table.style='Table Grid'
            for j,row in enumerate(vals):
                cells=table.add_row().cells
                for c,v in zip(cells,row):
                    c.text=v
                    for p in c.paragraphs:
                        p.paragraph_format.first_line_indent=0;p.paragraph_format.line_spacing=1.1
                        p.paragraph_format.space_after=Pt(4)
                        for run in p.runs:run.font.size=Pt(10);run.bold=j==0
            doc.add_paragraph();continue
        p=doc.add_paragraph(line)
        if line.startswith(('Fuente:','Figura.')):
            p.paragraph_format.first_line_indent=0;p.paragraph_format.line_spacing=1.1
            for run in p.runs:run.font.size=Pt(10)
        else:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for node in doc.element.iter(qn('w:pBdr')):node.getparent().remove(node)
    out=HERE/'objetivo3/CONTENIDO_OBJETIVO3.docx';doc.save(out)
    print(out)
if __name__=='__main__':build()
