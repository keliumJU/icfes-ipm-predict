"""Cotejo documental completo Total. No exporta panel ni incorpora cifras al análisis.
Exporta referencias de celdas y dictamen; valores solo en diferencias.
Compara números almacenados exactamente y por separado redondeo a un decimal.
"""
from inspeccionar import sheets, HERE, ROOT
from iniciar import sha
from openpyxl.utils import get_column_letter
from decimal import Decimal
import re,json,csv,collections
YEARS={2018,2019,2021,2022,2023,2024,2025}
def territory_name(x):
    # Rótulo Bogotá en privaciones frente a Bogotá D.C. en IPM. No cambio geográfico.
    return {'Bogotá':'Bogotá D.C.','Guainia':'Guainía','Vaupes':'Vaupés'}.get(str(x).strip(),str(x).strip())
def num(x): return isinstance(x,(int,float)) and not isinstance(x,bool)
def year(x):
    m=re.fullmatch(r'(20\d\d)(\*+)?',str(int(x)) if num(x) else str(x).strip())
    return int(m[1]) if m else None
def extract(p):
    book=dict(sheets(p)); ipm=book['IPM_Departamentos']
    varsheets=[n for n in book if re.fullmatch(r'IPM_(Variables|Indicadores)_Departamento\s*',n)]
    assert len(varsheets)==1,(p,varsheets)
    out={};structure=[];territories=[];variables=[]
    for name in ['IPM_Departamentos',varsheets[0]]:
        rows=book[name]; head=next(i for i,r in enumerate(rows) if i<20 and 'Total' in r)
        years={};last=None
        for c,v in enumerate(rows[head-1]):
            y=year(v)
            if y:last=y
            if rows[head][c]=='Total' and last in YEARS:years[c]=last
        assert years,(p,name)
        isvar=name!= 'IPM_Departamentos';territory=None; seen=collections.Counter()
        for i in range(head+1,len(rows)):
            r=rows[i]
            if not isvar:
                if not isinstance(r[0],str) or not any(num(r[c]) for c in years):continue
                territory=territory_name(r[0]); territories.append(territory);variable='IPM'
            else:
                if r[0] and territory_name(r[0]) in territories:territory=territory_name(r[0])
                if territory is None or not isinstance(r[1],str) or not any(num(r[c]) for c in years):continue
                variable=r[1].strip()
                # Later significance tables are not the rates table.
                if seen[territory]>=15:continue
                seen[territory]+=1
                if territory==territories[0]:variables.append(variable)
                assert variable in variables,(p,territory,variable)
            for c,y in years.items():
                assert num(r[c]),(p,name,i,c,r[c])
                key=(y,territory,variable);assert key not in out,key
                out[key]=(Decimal(str(r[c])),name,f'{get_column_letter(c+1)}{i+1}')
        notes=[dict(celda=f'{get_column_letter(c+1)}{i+1}',texto=str(v)) for i,r in enumerate(rows) for c,v in enumerate(r) if isinstance(v,str) and any(t in v.lower() for t in ['fuente:','cnpv','censo','publicación','actualizado','san andrés','representatividad','variable\n']) and (i<head or i>head+33)]
        structure.append(dict(hoja=name,encabezado_total=head+1,columnas_total={get_column_letter(c+1):y for c,y in years.items()},notas=notes))
    assert len(territories)==33,(p,len(territories))
    assert len(variables)==15,(p,variables)
    assert len(out)==33*16*len(years),(p,len(out),len(years))
    formats=set()
    if p.suffix=='.xlsx':
        from openpyxl import load_workbook
        raw=load_workbook(p,read_only=True,data_only=False)
        # Lectura por bloques completos; no recalcular ni modificar libro.
        for name in ['IPM_Departamentos',varsheets[0]]:
            wanted={loc for value,n,loc in out.values() if n==name}
            for cells in raw[name].iter_rows():
                for cell in cells:
                    if getattr(cell,'coordinate',None) in wanted:
                        assert num(cell.value), (p,name,cell.coordinate,'Fórmula o valor no numérico')
                        formats.add(cell.number_format)
        raw.close()
    else:
        import xlrd
        from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
        raw=xlrd.open_workbook(p,formatting_info=True)
        for value,name,loc in out.values():
            col,row=coordinate_from_string(loc);cell=raw.sheet_by_name(name).cell(row-1,column_index_from_string(col)-1)
            formats.add(raw.format_map[raw.xf_list[cell.xf_index].format_key].format_str)
    return out,dict(archivo=str(p.relative_to(ROOT)),sha256=sha(p),territorios=territories,variables=variables,estructura=structure,formatos_numericos_cotejados=sorted(formats),unidad='Porcentaje, sin dividir por 100',precision='Igualdad exacta de números almacenados; redondeo a un decimal separado',equivalencias_rotulos={'Bogotá':'Bogotá D.C.','Guainia':'Guainía','Vaupes':'Vaupés'})
def calculate():
    locals={y:ROOT/'Fuente del IPM'/f'anex-PMultidimensional-Departamental-{y}.xlsx' for y in [2024,2025]}
    official=sorted((HERE/'evidencia').glob('dane_*.xls*'))
    parsed={}; structures=[]
    for p in list(locals.values())+official:
        data,structure=extract(p);parsed[p]=data;structures.append(structure)
    evidence=[];summary=[];differences=[]
    for op in official:
        for ly,lp in locals.items():
            a,b=parsed[lp],parsed[op];common=sorted(set(a)&set(b))
            assert common
            counts=collections.defaultdict(lambda:collections.Counter())
            for key in common:
                y,t,v=key;av,an,ac=a[key];bv,bn,bc=b[key]
                eq=av==bv;rounded=av.quantize(Decimal('.1'))==bv.quantize(Decimal('.1'))
                row=dict(oficial=str(op.relative_to(HERE)),local=str(lp.relative_to(ROOT)),anio=y,territorio=t,variable=v,hoja_local=an,celda_local=ac,hoja_oficial=bn,celda_oficial=bc,igual_exacto=eq,igual_un_decimal=rounded)
                evidence.append(row);counts[y]['celdas']+=1;counts[y]['iguales']+=eq;counts[y]['diferentes']+=not eq
                if not eq:differences.append(dict(**row,valor_local=str(av),valor_oficial=str(bv),diferencia_pp=str(av-bv)))
            for y,c in counts.items():
                assert c['celdas']==528
                summary.append(dict(oficial=op.name,libro_local=ly,anio=y,**c,identidad_archivo=sha(lp)==sha(op)))
    return evidence,summary,differences,structures
def write_csv(name,rows):
    with (HERE/name).open('w',newline='') as f:
        cols=list(rows[0]) if rows else ['sin_diferencias']
        w=csv.DictWriter(f,fieldnames=cols);w.writeheader();w.writerows(rows)
if __name__=='__main__':
    e,s,d,st=calculate()
    write_csv('COTEJO_CELDAS.csv',e);write_csv('RESUMEN_COTEJO.csv',s);write_csv('DIFERENCIAS.csv',d)
    (HERE/'ESTRUCTURA_Y_NOTAS.json').write_text(json.dumps(st,ensure_ascii=False,indent=2)+'\n')
    (HERE/'RESUMEN_COTEJO.json').write_text(json.dumps(dict(comparaciones=len(e),iguales=sum(r['igual_exacto'] for r in e),diferentes=len(d),resumen=s,metodo='Total; unión por año/territorio/rótulo. Decimal de número almacenado; sin tolerancia para igualdad exacta. No exportación de panel.'),ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(comparaciones=len(e),diferentes=len(d),identidades=[(r['oficial'],r['libro_local']) for r in s if r['identidad_archivo']])))
